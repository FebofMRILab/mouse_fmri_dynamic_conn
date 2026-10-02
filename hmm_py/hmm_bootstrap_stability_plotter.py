#!/usr/bin/env python3
"""
hmm_bootstrap_stability_plotter.py   (v2)
-------------------------------------------------------------------
Supplementary figures and tables for the block-bootstrap stability
analysis of subject- and state-specific HMM FC networks.

Reads the CSV written by the updated MATLAB run5B (which includes the v2
diagnostics and a fit_id column). Older CSVs (MATLAB v1, Python v2) still
work: v2-only panels and rows are filled when their columns exist and are
otherwise omitted or labelled "not available".

Input
    bootstrap_stability_subject_state.csv   (one row per subject x state);
    default: <model folder>/hmm_bootstrap_stability/ of the reported fit
    (hmm_model_selection.py, same settings as hmm_config.m: HMM-PCA rank 15,
    K = 4). The file must carry the fit ID in <model folder>/fit_id.txt.

Outputs (in --out_dir)
    FigS_bootstrap_attrition.{pdf,png}   main supplementary figure:
        A retention by state, B group x state heatmap, C Neff vs FO,
        D bootstrap success, E stability metric comparison (point 1),
        F stability vs Neff with near-static states marked (points 2, 4)
    FigS_bootstrap_stability.{pdf,png}   Jaccard / rank rho detail
    FigS_bootstrap_occupancy.{pdf,png}   Neff, p/Neff, effective rank,
                                         subject max FO
    TableS_bootstrap_stability.docx      methods note, draft legend,
                                         Tables S-A..S-F
    TableS_*.csv                         same tables as CSV
    neff_covariates.csv                  subject x state Neff / log Neff /
                                         FO / near-static flag, to merge
                                         into the group analyses

Usage
    python hmm_bootstrap_stability_plotter.py                 # reported fit
    python hmm_bootstrap_stability_plotter.py --csv <subject_state.csv> \
        --out_dir <dir> [--min_neff 83]
"""

import argparse
import os
import warnings

import numpy as np
import pandas as pd
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.lines import Line2D
from matplotlib.patches import Patch

import hmm_model_selection as sel

# ------------------------------------------------------------------
# Settings (must match the analysis run)
# ------------------------------------------------------------------
MIN_NEFF = 83          # hmm_config.m min_effective_n (run5 and run5B); --min_neff
MIN_SUCCESS = 0.80
GRAPH_DENSITY = 0.15
N_BOOT = 1000
NEAR_STATIC_FO = 0.85
CI_LEVEL = "95%"

GROUP_ORDER_PREF = ["NTG_Y", "NTG_M", "NTG_A",
                    "CRND8_Y", "CRND8_M", "CRND8_A",
                    "5XFAD_Y", "5XFAD_M", "5XFAD_A"]
GENOTYPE_BASE = {"NTG": "#555555", "CRND8": "#0072B2", "5XFAD": "#D55E00"}
AGE_ALPHA = {"Y": 0.45, "M": 0.72, "A": 1.0}

LBL_OK = "Analyzable"
LBL_BOOT = "Bootstrap success < {:.0%}".format(MIN_SUCCESS)
LBL_NEFF = "Neff < {} (occupancy)".format(MIN_NEFF)
LBL_INV = "Invalid original network"
STATUS_COLORS = {LBL_OK: "#4C9A6A", LBL_BOOT: "#E0A526",
                 LBL_NEFF: "#B23A48", LBL_INV: "#7A5195"}
FIT_ID = "not recorded"


def set_min_neff(value):
    """Set the occupancy criterion and the labels that show it."""
    global MIN_NEFF, LBL_NEFF, STATUS_COLORS
    MIN_NEFF = value
    LBL_NEFF = "Neff < {} (occupancy)".format(MIN_NEFF)
    STATUS_COLORS = {LBL_OK: "#4C9A6A", LBL_BOOT: "#E0A526",
                     LBL_NEFF: "#B23A48", LBL_INV: "#7A5195"}

# Stability metrics: (Jaccard column, rho column, label, colour)
METRICS = [
    ("median_edge_jaccard", "median_nodal_strength_rank_correlation",
     "Resample vs full data", "#9E9E9E"),
    ("median_edge_jaccard_unfiltered", "median_nodal_rank_unfiltered",
     "Resample vs full data, unfiltered", "#D9D9D9"),
    ("median_edge_jaccard_rr", "median_nodal_rank_rr",
     "Resample vs resample", "#6BAED6"),
    ("split_half_edge_jaccard", "split_half_nodal_rank_correlation",
     "Split half", "#2171B5"),
]
STATE_MARKERS = ["o", "s", "^", "D", "v", "P"]

plt.rcParams.update({
    "font.family": "sans-serif",
    "font.sans-serif": ["Arial", "Helvetica", "DejaVu Sans"],
    "font.size": 8, "axes.titlesize": 9, "axes.labelsize": 8,
    "xtick.labelsize": 7, "ytick.labelsize": 7, "legend.fontsize": 7,
    "axes.spines.top": False, "axes.spines.right": False,
    "axes.linewidth": 0.8, "pdf.fonttype": 42, "ps.fonttype": 42,
    "savefig.dpi": 600,
})


# ------------------------------------------------------------------
# Helpers
# ------------------------------------------------------------------
def has(df, col):
    return col in df.columns and df[col].notna().any()


def group_order(groups):
    present = list(pd.unique(groups))
    ordered = [g for g in GROUP_ORDER_PREF if g in present]
    return ordered + sorted(g for g in present if g not in ordered)


def group_color(g):
    geno, _, age = g.partition("_")
    base = matplotlib.colors.to_rgb(GENOTYPE_BASE.get(geno, "#888888"))
    a = AGE_ALPHA.get(age, 0.8)
    return tuple(1 - a * (1 - c) for c in base)


def classify(row):
    if row["analyzable"]:
        return LBL_OK
    if not row["valid_original_network"]:
        return LBL_INV
    if not row["passes_Neff"]:
        return LBL_NEFF
    return LBL_BOOT


def jaccard_chance(d):
    """Expected Jaccard of two independent random edge sets of density d."""
    return d / (2.0 - d)


def med_iqr(x, fmt="{:.2f}"):
    x = pd.Series(x).dropna()
    if x.empty:
        return "—"
    q1, q2, q3 = np.percentile(x, [25, 50, 75])
    return (fmt + " [" + fmt + "–" + fmt + "]").format(q2, q1, q3)


def save(fig, out_dir, stem):
    for ext in ("pdf", "png"):
        fig.savefig(os.path.join(out_dir, f"{stem}.{ext}"),
                    bbox_inches="tight")
    plt.close(fig)


def neff_log_axis(ax, axis="x"):
    """Log axis with plain tick labels (10, 20, 50, ...)."""
    from matplotlib.ticker import FixedLocator, NullFormatter, ScalarFormatter
    a = ax.xaxis if axis == "x" else ax.yaxis
    (ax.set_xscale if axis == "x" else ax.set_yscale)("log")
    a.set_major_locator(FixedLocator([5, 10, 20, 50, 100, 200, 500, 1000]))
    fmt = ScalarFormatter()
    fmt.set_scientific(False)
    a.set_major_formatter(fmt)
    a.set_minor_formatter(NullFormatter())


def panel_label(ax, s):
    ax.text(-0.22, 1.10, s, transform=ax.transAxes, fontsize=11,
            fontweight="bold", va="bottom", ha="left")


def state_ticks(ax, states, df=None):
    ax.set_xticks(states)
    if df is None:
        ax.set_xticklabels([f"State {k}" for k in states])
    else:
        ax.set_xticklabels([f"State {k}\n(n={int((df.state == k).sum())})"
                            for k in states])
    ax.set_xlim(min(states) - 0.6, max(states) + 0.6)


def strip_by_group(ax, df, ycol, states, groups, jitter=0.28, s=10,
                   show_box=True, ring_near_static=False):
    rng = np.random.default_rng(1)
    if show_box:
        data = [df.loc[df.state == k, ycol].dropna().values for k in states]
        ax.boxplot(data, positions=states, widths=0.55, showfliers=False,
                   patch_artist=True,
                   medianprops=dict(color="black", lw=1.2),
                   boxprops=dict(facecolor="#EEEEEE", lw=0.8),
                   whiskerprops=dict(lw=0.8), capprops=dict(lw=0.8))
    for k in states:
        sub = df[df.state == k]
        for g in groups:
            sg = sub[(sub.group == g) & sub[ycol].notna()]
            x = k + rng.uniform(-jitter, jitter, size=len(sg))
            ax.scatter(x, sg[ycol], s=s, color=group_color(g),
                       edgecolor="white", linewidth=0.3, zorder=3)
            if ring_near_static:
                ns = sg.near_static.values
                ax.scatter(x[ns], sg[ycol].values[ns], s=s * 3.5,
                           facecolor="none", edgecolor="black",
                           linewidth=0.8, zorder=4)
    state_ticks(ax, states)


def group_handles(groups):
    return [Line2D([], [], marker="o", ls="", markersize=5,
                   markerfacecolor=group_color(g), markeredgecolor="white",
                   label=g) for g in groups]


def state_handles(states):
    return [Line2D([], [], marker=STATE_MARKERS[i], ls="", color="grey",
                   markersize=4, label=f"State {k}")
            for i, k in enumerate(states)]


NEAR_STATIC_HANDLE = Line2D([], [], marker="o", ls="", markersize=7,
                            markerfacecolor="none", markeredgecolor="black",
                            label=f"Near-static (FO ≥ {NEAR_STATIC_FO})")


def available_metrics(df, which):
    """which = 0 for Jaccard, 1 for rho."""
    return [(m[which], m[2], m[3]) for m in METRICS if has(df, m[which])]


def metric_boxes(ax, df, metrics, states):
    """Grouped boxes: one box per available metric within each state."""
    n = len(metrics)
    width = 0.7 / n
    for i, (col, lab, colr) in enumerate(metrics):
        pos = [k - 0.35 + width * (i + 0.5) for k in states]
        data = [df.loc[df.state == k, col].dropna().values for k in states]
        data = [d if d.size else np.array([np.nan]) for d in data]
        ax.boxplot(data, positions=pos, widths=width * 0.85,
                   showfliers=False, patch_artist=True,
                   medianprops=dict(color="black", lw=1.0),
                   boxprops=dict(facecolor=colr, lw=0.6),
                   whiskerprops=dict(lw=0.6), capprops=dict(lw=0.6))
    state_ticks(ax, states)
    return [Patch(facecolor=c, edgecolor="black", lw=0.6, label=l)
            for _, l, c in metrics]


def vs_neff(ax, df, ycol, lo, hi, states, extra_cols=(),
            primary_label="Resample vs full data"):
    """Stability vs Neff: per subject-state median and CI, near-static
    ringed, Spearman rho with Neff annotated for each available metric."""
    for i, k in enumerate(states):
        s = df[(df.state == k) & df[ycol].notna()]
        if has(s, lo) and has(s, hi):
            ax.errorbar(s.Neff, s[ycol],
                        yerr=[s[ycol] - s[lo], s[hi] - s[ycol]],
                        fmt="none", ecolor="#CCCCCC", elinewidth=0.5,
                        zorder=1)
        ax.scatter(s.Neff, s[ycol], s=12, marker=STATE_MARKERS[i],
                   c=[group_color(g) for g in s.group], edgecolor="white",
                   linewidth=0.3, zorder=3)
        ns = s[s.near_static]
        ax.scatter(ns.Neff, ns[ycol], s=45, facecolor="none",
                   edgecolor="black", linewidth=0.8, zorder=4)
    lines = []
    for col, lab in [(ycol, primary_label)] + list(extra_cols):
        if not has(df, col):
            continue
        d = df[["Neff", col]].dropna()
        rho = d.corr(method="spearman").iat[0, 1]
        lines.append(f"{lab}: ρ = {rho:.2f}")
    ax.text(0.03, 0.03, "Spearman ρ with Neff\n" + "\n".join(lines),
            transform=ax.transAxes, va="bottom", fontsize=6.5, zorder=5,
            bbox=dict(facecolor="white", edgecolor="none", alpha=0.85,
                      pad=1.5))
    neff_log_axis(ax, "x")
    ax.set_xlabel("Effective sample size, Neff (log)")


# ------------------------------------------------------------------
# Figure: attrition (main supplementary figure)
# ------------------------------------------------------------------
def fig_attrition(df, states, groups, out_dir):
    fig = plt.figure(figsize=(7.2, 9.6))
    gs = fig.add_gridspec(3, 2, hspace=0.62, wspace=0.40)

    # A. retention by state
    ax = fig.add_subplot(gs[0, 0])
    status_order = [s for s in STATUS_COLORS if s in df.status.unique()]
    counts = (df.groupby(["state", "status"]).size().unstack(fill_value=0)
                .reindex(index=states, columns=status_order, fill_value=0))
    bottom = np.zeros(len(states))
    for st in status_order:
        v = counts[st].values
        ax.bar(states, v, bottom=bottom, color=STATUS_COLORS[st], width=0.6,
               edgecolor="white", lw=0.5, label=st)
        for x, b, n in zip(states, bottom, v):
            if n > 0 and st != LBL_OK:
                ax.text(x + 0.34, b + n / 2, str(n), va="center",
                        ha="left", fontsize=7, color=STATUS_COLORS[st])
        bottom += v
    totals = counts.sum(axis=1).values
    kept = counts[LBL_OK].values
    for x, t, k in zip(states, totals, kept):
        ax.text(x, t + 1.2, f"{k}/{t}\n({100 * k / t:.1f}%)", ha="center",
                va="bottom", fontsize=7)
    ax.set_xticks(states)
    ax.set_xticklabels([f"State {k}" for k in states])
    ax.set_ylabel("Subject-states (n)")
    ax.set_ylim(0, totals.max() * 1.22)
    ax.set_xlim(min(states) - 0.6, max(states) + 0.8)
    ax.legend(loc="upper center", bbox_to_anchor=(0.5, -0.12), ncol=2,
              frameon=False, handlelength=1.0)
    ax.set_title("Retention by state")
    panel_label(ax, "A")

    # B. heatmap group x state
    ax = fig.add_subplot(gs[0, 1])
    tab_k = (df.groupby(["group", "state"]).analyzable.sum().unstack()
               .reindex(index=groups, columns=states))
    tab_n = (df.groupby(["group", "state"]).size().unstack()
               .reindex(index=groups, columns=states))
    pct = 100 * tab_k / tab_n
    # 0-100%: with Neff >= 83 many cells fall below 50%, which a 50-100% scale
    # would draw identically (white)
    im = ax.imshow(pct.values, cmap="Greens", vmin=0, vmax=100,
                   aspect="auto")
    for i in range(len(groups)):
        for j in range(len(states)):
            k, n, p = tab_k.iat[i, j], tab_n.iat[i, j], pct.iat[i, j]
            ax.text(j, i, f"{int(k)}/{int(n)}", ha="center", va="center",
                    fontsize=7, color="white" if p > 85 else "black")
    ax.set_xticks(range(len(states)))
    ax.set_xticklabels([f"State {k}" for k in states])
    ax.set_yticks(range(len(groups)))
    ax.set_yticklabels(groups)
    ax.spines[:].set_visible(False)
    ax.tick_params(length=0)
    cb = fig.colorbar(im, ax=ax, fraction=0.046, pad=0.03)
    cb.set_label("Analyzable (%)")
    cb.outline.set_linewidth(0.5)
    ax.set_title("Analyzable subject-states by group")
    panel_label(ax, "B")

    # C. Neff vs FO
    ax = fig.add_subplot(gs[1, 0])
    for st in status_order:
        sub = df[df.status == st]
        for i, k in enumerate(states):
            s2 = sub[sub.state == k]
            ax.scatter(s2.FO, s2.Neff, s=12, marker=STATE_MARKERS[i],
                       color=STATUS_COLORS[st], edgecolor="white",
                       linewidth=0.3, zorder=3 if st != LBL_OK else 2)
    ns = df[df.near_static]
    ax.scatter(ns.FO, ns.Neff, s=45, facecolor="none", edgecolor="black",
               linewidth=0.8, zorder=4)
    ax.axhline(MIN_NEFF, color="black", ls="--", lw=0.8)
    ax.axvline(NEAR_STATIC_FO, color="grey", ls=":", lw=0.8)
    ax.text(0.99, MIN_NEFF * 1.08, f"Neff = {MIN_NEFF}",
            transform=ax.get_yaxis_transform(), ha="right", va="bottom",
            fontsize=7)
    neff_log_axis(ax, "y")
    ax.set_xlabel("Fractional occupancy (FO)")
    ax.set_ylabel("Effective sample size, Neff (log)")
    ax.set_xlim(-0.02, 1.02)
    ax.legend(handles=state_handles(states) + [NEAR_STATIC_HANDLE],
              loc="best", frameon=False)
    ax.set_title("Occupancy criterion")
    panel_label(ax, "C")

    # D. bootstrap success
    ax = fig.add_subplot(gs[1, 1])
    boot = df[df.passes_Neff & df.valid_original_network]
    strip_by_group(ax, boot, "bootstrap_success_fraction", states, groups,
                   show_box=False)
    ax.axhline(MIN_SUCCESS, color="black", ls="--", lw=0.8)
    ax.text(0.99, MIN_SUCCESS - 0.02, f"{MIN_SUCCESS:.0%} criterion",
            transform=ax.get_yaxis_transform(), ha="right", va="top",
            fontsize=7)
    ax.set_ylabel(f"Successful resamples\n(fraction of {N_BOOT})")
    ax.set_ylim(min(0.0, boot.bootstrap_success_fraction.min() - 0.05),
                1.04)
    ax.set_title("Bootstrap success (Neff-passing only)")
    ax.legend(handles=group_handles(groups), loc="lower left",
              frameon=False, ncol=2, handletextpad=0.2, columnspacing=0.6)
    panel_label(ax, "D")

    ana = df[df.analyzable]

    # E. stability metric comparison (Jaccard)
    ax = fig.add_subplot(gs[2, 0])
    mets = available_metrics(ana, 0)
    handles = metric_boxes(ax, ana, mets, states)
    chance = jaccard_chance(GRAPH_DENSITY)
    ax.axhline(chance, color="grey", ls=":", lw=0.9)
    ax.text(0.99, chance + 0.01, f"chance ≈ {chance:.3f}",
            transform=ax.get_yaxis_transform(), ha="right", va="bottom",
            fontsize=6.5, color="grey")
    ax.set_ylim(0, 1)
    ax.set_ylabel("Edge-set Jaccard")
    ax.set_title("Edge-set stability by comparison")
    ax.legend(handles=handles, loc="upper center",
              bbox_to_anchor=(0.5, -0.12), ncol=2, frameon=False,
              handlelength=1.0)
    if len(mets) == 1:
        ax.text(0.5, 0.92, "v2 metrics not in this CSV",
                transform=ax.transAxes, ha="center", fontsize=7,
                color="grey")
    panel_label(ax, "E")

    # F. stability vs Neff
    ax = fig.add_subplot(gs[2, 1])
    extra = [(c, l) for c, l, _ in mets[1:]
             if c in ("median_edge_jaccard_rr", "split_half_edge_jaccard")]
    vs_neff(ax, ana, "median_edge_jaccard", "edge_jaccard_CI_low",
            "edge_jaccard_CI_high", states,
            extra_cols=extra)
    ax.set_ylabel(f"Edge-set Jaccard vs full data\n"
                  f"(median, {CI_LEVEL} interval)")
    ax.set_ylim(0, 1)
    ax.set_title("Stability depends on occupancy")
    # below the axis, as in panel E, so it cannot cover the Spearman box
    ax.legend(handles=state_handles(states) + [NEAR_STATIC_HANDLE],
              loc="upper center", bbox_to_anchor=(0.5, -0.24), ncol=2,
              frameon=False, handlelength=1.0)
    panel_label(ax, "F")

    save(fig, out_dir, "FigS_bootstrap_attrition")


# ------------------------------------------------------------------
# Figure: stability detail
# ------------------------------------------------------------------
def fig_stability(df, states, groups, out_dir):
    ana = df[df.analyzable]
    chance = jaccard_chance(GRAPH_DENSITY)
    fig, axes = plt.subplots(2, 2, figsize=(7.2, 6.4),
                             gridspec_kw=dict(hspace=0.55, wspace=0.38))

    ax = axes[0, 0]
    strip_by_group(ax, ana, "median_edge_jaccard", states, groups,
                   ring_near_static=True)
    ax.axhline(chance, color="grey", ls=":", lw=0.9)
    ax.text(0.99, chance + 0.01,
            f"chance ≈ {chance:.3f} (density {GRAPH_DENSITY:.0%})",
            transform=ax.get_yaxis_transform(), ha="right", va="bottom",
            fontsize=6.5, color="grey")
    ax.set_ylim(0, 1)
    ax.set_ylabel("Edge-set Jaccard vs full data")
    ax.set_title("Edge-set stability by group")
    state_ticks(ax, states, ana)
    panel_label(ax, "A")

    ax = axes[0, 1]
    strip_by_group(ax, ana, "median_nodal_strength_rank_correlation",
                   states, groups, ring_near_static=True)
    ax.axhline(0, color="grey", ls=":", lw=0.9)
    ax.set_ylim(min(-0.05, ana.median_nodal_strength_rank_correlation.min()
                    - 0.05), 1)
    ax.set_ylabel("Nodal-strength Spearman ρ vs full data")
    ax.set_title("Nodal-strength rank stability by group")
    state_ticks(ax, states, ana)
    ax.legend(handles=group_handles(groups) + [NEAR_STATIC_HANDLE],
              loc="lower right", frameon=False, ncol=2, handletextpad=0.2,
              columnspacing=0.6)
    panel_label(ax, "B")

    ax = axes[1, 0]
    mets = available_metrics(ana, 1)
    handles = metric_boxes(ax, ana, mets, states)
    ax.axhline(0, color="grey", ls=":", lw=0.9)
    ax.set_ylabel("Nodal-strength Spearman ρ")
    ax.set_title("Rank stability by comparison")
    ax.legend(handles=handles, loc="upper center",
              bbox_to_anchor=(0.5, -0.12), ncol=2, frameon=False,
              handlelength=1.0)
    panel_label(ax, "C")

    ax = axes[1, 1]
    extra = [(c, l) for c, l, _ in mets[1:]
             if c in ("median_nodal_rank_rr",
                      "split_half_nodal_rank_correlation")]
    vs_neff(ax, ana, "median_nodal_strength_rank_correlation",
            "nodal_rank_correlation_CI_low",
            "nodal_rank_correlation_CI_high", states, extra_cols=extra)
    ax.set_ylabel(f"Nodal-strength ρ vs full data\n"
                  f"(median, {CI_LEVEL} interval)")
    ax.set_ylim(None, 1)
    ax.set_title("Rank stability vs Neff")
    ax.legend(handles=state_handles(states), loc="lower right",
              frameon=False)
    panel_label(ax, "D")

    save(fig, out_dir, "FigS_bootstrap_stability")


# ------------------------------------------------------------------
# Figure: occupancy / conditioning
# ------------------------------------------------------------------
def fig_occupancy(df, states, groups, out_dir, p):
    fig, axes = plt.subplots(2, 2, figsize=(7.2, 5.8),
                             gridspec_kw=dict(hspace=0.55, wspace=0.45))

    ax = axes[0, 0]
    strip_by_group(ax, df, "Neff", states, groups, ring_near_static=True)
    ax.axhline(MIN_NEFF, color="black", ls="--", lw=0.8)
    neff_log_axis(ax, "y")
    ax.set_ylabel("Neff (log)")
    ax.set_title("Effective sample size")
    panel_label(ax, "A")

    ax = axes[0, 1]
    strip_by_group(ax, df, "p_over_Neff", states, groups)
    ax.axhline(1.0, color="grey", ls=":", lw=0.9)
    ax.axhline(p / MIN_NEFF, color="black", ls="--", lw=0.8)
    ax.set_yscale("log")
    ax.set_ylabel("p / Neff (log)")
    ax.set_title(f"p / Neff (p = {p})")
    panel_label(ax, "B")

    ax = axes[1, 0]
    strip_by_group(ax, df, "covariance_effective_rank", states, groups)
    ax.set_ylabel("Entropy effective rank")
    ax.set_title("Covariance effective rank")
    panel_label(ax, "C")

    # D. subject-level dominance
    ax = axes[1, 1]
    subj = (df.groupby(["subject_id", "group"])
              .agg(max_FO=("FO", "max")).reset_index())
    rng = np.random.default_rng(2)
    for i, g in enumerate(groups):
        y = subj.loc[subj.group == g, "max_FO"].values
        x = i + rng.uniform(-0.25, 0.25, size=y.size)
        ax.scatter(x, y, s=12, color=group_color(g), edgecolor="white",
                   linewidth=0.3, zorder=3)
        ns = y >= NEAR_STATIC_FO
        ax.scatter(x[ns], y[ns], s=45, facecolor="none", edgecolor="black",
                   linewidth=0.8, zorder=4)
    ax.axhline(NEAR_STATIC_FO, color="black", ls="--", lw=0.8)
    ax.axhline(1 / len(states), color="grey", ls=":", lw=0.8)
    ax.set_xticks(range(len(groups)))
    ax.set_xticklabels(groups, rotation=45, ha="right")
    ax.set_ylim(0, 1.04)
    ax.set_ylabel("Max FO across states (per subject)")
    ax.set_title("State dominance per subject")
    ax.legend(handles=group_handles(groups), loc="upper left",
              bbox_to_anchor=(1.0, 1.0), frameon=False)
    panel_label(ax, "D")

    save(fig, out_dir, "FigS_bootstrap_occupancy")


# ------------------------------------------------------------------
# Neff-adjusted group model (point 2)
# ------------------------------------------------------------------
def group_models(df, groups):
    """OLS of each stability metric on group + state, with and without
    log Neff, cluster-robust SEs by subject (states nested in subjects).
    Joint Wald test for the group terms."""
    try:
        import statsmodels.formula.api as smf
    except ImportError:
        return None

    ref = groups[0]
    targets = [("median_edge_jaccard", "Edge Jaccard vs full data"),
               ("median_nodal_strength_rank_correlation",
                "Nodal rank ρ vs full data"),
               ("median_edge_jaccard_rr", "Edge Jaccard, resample vs resample"),
               ("split_half_edge_jaccard", "Edge Jaccard, split half")]
    rows = []
    for col, lab in targets:
        if not has(df, col):
            continue
        base = df[df.analyzable & df[col].notna()].copy()
        base["y"] = base[col]
        base["log_Neff"] = np.log(base.Neff)
        specs = [("group + state", base, False),
                 ("group + state + log Neff", base, True),
                 ("… excluding near-static", base[~base.near_static], True)]
        for name, d, with_neff in specs:
            if d.group.nunique() < 2 or len(d) < 10:
                continue
            f = (f"y ~ C(group, Treatment('{ref}')) + C(state)"
                 + (" + log_Neff" if with_neff else ""))
            with warnings.catch_warnings():
                warnings.simplefilter("ignore")
                res = smf.ols(f, data=d).fit(
                    cov_type="cluster",
                    cov_kwds={"groups": d.subject_id.astype(int)})
            names = list(res.params.index)
            gidx = [i for i, n in enumerate(names) if n.startswith("C(group")]
            L = np.zeros((len(gidx), len(names)))
            for r_, i in enumerate(gidx):
                L[r_, i] = 1
            w = res.wald_test(L, scalar=True)
            stat = float(np.squeeze(w.statistic))
            row = {"Metric": lab, "Model": name, "n": int(res.nobs),
                   "Group Wald χ² (df)": f"{stat:.2f} ({len(gidx)})",
                   "Group p": f"{float(w.pvalue):.3g}"}
            if with_neff:
                b = res.params["log_Neff"]
                lo, hi = res.conf_int().loc["log_Neff"]
                row["β log Neff [95% CI]"] = f"{b:.3f} [{lo:.3f}, {hi:.3f}]"
                row["log Neff p"] = f"{res.pvalues['log_Neff']:.3g}"
            else:
                row["β log Neff [95% CI]"] = "—"
                row["log Neff p"] = "—"
            rows.append(row)
    return pd.DataFrame(rows)


# ------------------------------------------------------------------
# Tables
# ------------------------------------------------------------------
def build_tables(df, states, groups):
    # A. state summary
    rows = []
    for k in states:
        s = df[df.state == k]
        a = s[s.analyzable]
        r = {"State": k, "N": len(s),
             "Neff ≥ {} n (%)".format(MIN_NEFF):
                 f"{int(s.passes_Neff.sum())} ({100 * s.passes_Neff.mean():.1f})",
             "Analyzable n (%)":
                 f"{int(s.analyzable.sum())} ({100 * s.analyzable.mean():.1f})",
             "Near-static n": int(s.near_static.sum()),
             "FO": med_iqr(s.FO), "Neff": med_iqr(s.Neff, "{:.0f}"),
             "p/Neff": med_iqr(s.p_over_Neff),
             "Eff. rank": med_iqr(s.covariance_effective_rank, "{:.1f}")}
        for jc, rc, lab, _ in METRICS:
            if has(df, jc):
                r[f"Jaccard: {lab}†"] = med_iqr(a[jc])
        for jc, rc, lab, _ in METRICS:
            if has(df, rc):
                r[f"Rank ρ: {lab}†"] = med_iqr(a[rc])
        rows.append(r)
    tA = pd.DataFrame(rows).set_index("State").T.reset_index()
    tA.columns = ["Measure"] + [f"State {k}" for k in states]

    # B. retention by group x state
    rows = []
    for g in groups + ["Total"]:
        s = df if g == "Total" else df[df.group == g]
        r = {"Group": g, "Subjects": int(s.subject_id.nunique())}
        for k in states:
            s2 = s[s.state == k]
            r[f"State {k}"] = (f"{int(s2.analyzable.sum())}/{len(s2)} "
                               f"({100 * s2.analyzable.mean():.0f}%)")
        per = s.groupby("subject_id").analyzable.sum()
        r["All states analyzable"] = \
            f"{int((per == len(states)).sum())}/{per.size}"
        r["Near-static subjects"] = int(
            s.groupby("subject_id").near_static.any().sum())
        rows.append(r)
    tB = pd.DataFrame(rows)

    # C. excluded subject-states
    gpos = {g: i for i, g in enumerate(groups)}
    ex = df[~df.analyzable].copy()
    ex["_g"] = ex.group.map(gpos)
    ex = ex.sort_values(["_g", "subject_id", "state"])
    tC = pd.DataFrame({
        "Subject": ex.subject_id.astype(int), "Group": ex.group,
        "State": ex.state.astype(int), "FO": ex.FO.map("{:.3f}".format),
        "Neff": ex.Neff.map("{:.1f}".format),
        "Successful resamples": np.where(
            ex.passes_Neff,
            ex.bootstrap_success_count.astype(int).astype(str)
            + f"/{N_BOOT}", "not run"),
        "Reason": ex.status,
    })
    if has(df, "n_fail_boot_Neff"):
        tC["Resample failures (Neff / network)"] = np.where(
            ex.passes_Neff,
            ex.n_fail_boot_Neff.fillna(0).astype(int).astype(str) + " / "
            + ex.n_fail_boot_network.fillna(0).astype(int).astype(str), "—")
    dom = (df.groupby("subject_id").FO.max() >= NEAR_STATIC_FO)
    tC["Subject has near-static state"] = np.where(
        ex.subject_id.map(dom), "yes", "")

    # D. stability by group x state
    rows = []
    for g in groups:
        for k in states:
            a = df[(df.group == g) & (df.state == k) & df.analyzable]
            r = {"Group": g, "State": k, "n": len(a),
                 "Neff": med_iqr(a.Neff, "{:.0f}"),
                 "Jaccard vs full": med_iqr(a.median_edge_jaccard),
                 "Rank ρ vs full":
                     med_iqr(a.median_nodal_strength_rank_correlation)}
            if has(df, "median_edge_jaccard_rr"):
                r["Jaccard resample-resample"] = \
                    med_iqr(a.median_edge_jaccard_rr)
            if has(df, "split_half_edge_jaccard"):
                r["Jaccard split half"] = med_iqr(a.split_half_edge_jaccard)
            rows.append(r)
    tD = pd.DataFrame(rows)

    # E. group models
    tE = group_models(df, groups)

    # F. near-static subject-states
    ns = df[df.near_static].copy()
    ns["_g"] = ns.group.map(gpos)
    ns = ns.sort_values(["_g", "subject_id", "state"])
    # map_count: Viterbi labels in the updated run5B (fit_id present);
    # most-probable-state labels in older CSVs
    count_label = "Viterbi volumes" if "fit_id" in df.columns else "MAP volumes"
    tF = pd.DataFrame({
        "Subject": ns.subject_id.astype(int), "Group": ns.group,
        "State": ns.state.astype(int), "FO": ns.FO.map("{:.3f}".format),
        count_label: ns.map_count.astype(int),
        "Analyzable": np.where(ns.analyzable, "yes", "no"),
        "Jaccard vs full": ns.median_edge_jaccard.map(
            lambda v: "—" if pd.isna(v) else f"{v:.2f}"),
    })
    if has(df, "subject_n_transitions"):
        tF["Transitions (subject)"] = ns.subject_n_transitions.astype(int)
        dwell_col = "mean_dwell_TR" if "mean_dwell_TR" in ns.columns else "mean_dwell_volumes"
        tF["Mean dwell (vol)"] = ns[dwell_col].map(
            lambda v: "—" if pd.isna(v) else f"{v:.0f}")
    other = (df[~df.near_static & df.subject_id.isin(ns.subject_id)]
             .groupby("subject_id").analyzable.apply(lambda x: (~x).sum()))
    tF["Other states excluded"] = ns.subject_id.map(other).fillna(0) \
        .astype(int).values
    return tA, tB, tC, tD, tE, tF


def legend_text(df, states):
    tot = df.groupby("state").size()
    kept = df.groupby("state").analyzable.sum()
    ret = ", ".join(f"{int(kept[k])}/{int(tot[k])} "
                    f"({100 * kept[k] / tot[k]:.1f}%)" for k in states)
    per = df.groupby("subject_id").analyzable.sum()
    n_all = int((per == len(states)).sum())
    n_ns = int(df.groupby("subject_id").near_static.any().sum())
    v2 = has(df, "median_edge_jaccard_rr")
    e_txt = ("Edge-set Jaccard for analyzable subject-states, computed as "
             "the similarity of each resample to the full-data network "
             "(grey), without the resample-level occupancy filter (light "
             "grey), between disjoint pairs of independent resamples "
             "(light blue), and between networks estimated from the first "
             "and second halves of the run (dark blue); the dotted line is "
             "the Jaccard expected for two random graphs at the same "
             "density."
             if v2 else
             "Edge-set Jaccard similarity of each resample to the full-data "
             "network for analyzable subject-states; the dotted line is the "
             "Jaccard expected for two random graphs at the same density.")
    return (
        "Figure S#. Retention and stability of subject-specific, "
        "state-specific networks in the block-bootstrap analysis. "
        "(A) Subject-states per HMM state that were analyzable (green) or "
        f"excluded, by reason. Subject-states with Kish effective sample "
        f"size of posterior state weights (Neff) below {MIN_NEFF} were not "
        f"bootstrapped (red); subject-states with fewer than "
        f"{MIN_SUCCESS:.0%} of {N_BOOT} circular moving-block resamples "
        f"yielding a valid network were excluded (orange). Resamples whose "
        f"own Neff fell below {MIN_NEFF} were counted as unsuccessful. "
        f"Retention for States {', '.join(map(str, states))}: {ret}. "
        "(B) Analyzable/total subject-states by group and state. "
        "(C) Neff against fractional occupancy; marker shape denotes state, "
        f"dashed line the Neff criterion, and circled points states "
        f"occupying ≥ {NEAR_STATIC_FO:.0%} of the scan (near-static). "
        f"(D) Fraction of successful resamples among subject-states meeting "
        f"the occupancy criterion. (E) {e_txt} "
        "(F) Edge-set Jaccard against Neff for analyzable subject-states, "
        f"with {CI_LEVEL} bootstrap percentile intervals; near-static states "
        f"are circled. {n_all} of {df.subject_id.nunique()} subjects "
        f"contributed an analyzable network for all states; {n_ns} "
        "subjects had a near-static state.")


def methods_note(df):
    v2 = has(df, "median_edge_jaccard_rr")
    txt = [
        "Interpretation notes. (i) Similarity of a resample to the "
        "full-data network overstates reliability, because a circular block "
        "bootstrap resample contains, on average, about 63% of the unique "
        "blocks in the original run. "
        + ("We therefore also report similarity between disjoint pairs of "
           "independent resamples, and between networks estimated from "
           "non-overlapping halves of the run. The halves share no time "
           "points, although the posterior probabilities come from an HMM "
           "fitted to the whole run. "
           if v2 else
           "Resample-vs-resample and split-half estimates are reported by "
           "the v2 analysis. "),
        "(ii) Resamples whose own Neff fell below the occupancy criterion "
        "were treated as unsuccessful, which conditions the stability "
        "estimates on adequate resampled occupancy. "
        + ("Unfiltered estimates, which retain these resamples, are shown "
           "for comparison. " if v2 else ""),
        "(iii) Stability increased with Neff (panel F and Table S-E), so "
        "group differences in state-specific connectivity can partly "
        "reflect group differences in estimation precision. Table S-E "
        "reports group effects on stability with and without log Neff as "
        "a covariate, and after excluding near-static states. "
        "(iv) Subjects who occupied one state for most of the scan "
        f"(FO ≥ {NEAR_STATIC_FO:.2f}) contribute a network for that state "
        "that approximates static functional connectivity. Their other "
        "states usually fail the occupancy criterion (Tables S-C and S-F).",
    ]
    return " ".join(txt)


def write_docx(tables, path, df, p):
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, Inches

    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Inches(11), Inches(8.5)
    for m in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, m, Inches(0.7))
    st = doc.styles["Normal"]
    st.font.name = "Arial"
    st.font.size = Pt(9)
    st.paragraph_format.space_after = Pt(0)
    st.paragraph_format.space_before = Pt(0)

    def heading(text):
        h = doc.add_paragraph()
        r = h.add_run(text)
        r.bold = True
        r.font.size = Pt(10)
        h.paragraph_format.space_after = Pt(4)

    def para(text, italic=False, size=8.5):
        n = doc.add_paragraph()
        rn = n.add_run(text)
        rn.italic = italic
        rn.font.size = Pt(size)
        n.paragraph_format.space_before = Pt(3)
        doc.add_paragraph()

    def set_borders(tbl):
        tblPr = tbl._tbl.tblPr
        b = OxmlElement("w:tblBorders")
        for edge in ("top", "bottom"):
            e = OxmlElement(f"w:{edge}")
            e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "8")
            e.set(qn("w:color"), "000000")
            b.append(e)
        for edge in ("left", "right", "insideH", "insideV"):
            e = OxmlElement(f"w:{edge}")
            e.set(qn("w:val"), "nil")
            b.append(e)
        tblPr.append(b)
        for cell in tbl.rows[0].cells:
            tcPr = cell._tc.get_or_add_tcPr()
            bd = OxmlElement("w:tcBorders")
            e = OxmlElement("w:bottom")
            e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "6")
            e.set(qn("w:color"), "000000")
            bd.append(e)
            tcPr.append(bd)

    def add_table(title, df_t, note):
        heading(title)
        if df_t is None or df_t.empty:
            para("Not available (requires v2 columns or statsmodels).",
                 italic=True)
            return
        t = doc.add_table(rows=1, cols=df_t.shape[1])
        for j, c in enumerate(df_t.columns):
            cell = t.rows[0].cells[j]
            cell.text = ""
            run = cell.paragraphs[0].add_run(str(c))
            run.bold = True
            run.font.size = Pt(8)
        for _, row in df_t.iterrows():
            cells = t.add_row().cells
            for j, v in enumerate(row):
                cells[j].text = ""
                run = cells[j].paragraphs[0].add_run(str(v))
                run.font.size = Pt(8)
        for row in t.rows:
            for j, cell in enumerate(row.cells):
                cell.paragraphs[0].alignment = (
                    WD_ALIGN_PARAGRAPH.LEFT if j == 0
                    else WD_ALIGN_PARAGRAPH.CENTER)
        set_borders(t)
        para(note, italic=True, size=8)

    tA, tB, tC, tD, tE, tF = tables
    T = int(df.n_timepoints.iloc[0]) if df.n_timepoints.nunique() == 1 else None
    L = int(df.block_length.iloc[0]) if df.block_length.nunique() == 1 else None
    common = (f"Circular moving-block bootstrap ({N_BOOT} resamples"
              + (f", block length {L} volumes, T = {T}" if L and T else "")
              + f") of paired time series and HMM posterior probabilities; "
              f"state-weighted correlation among p = {p} ROIs; positive-edge "
              f"proportional threshold retaining {GRAPH_DENSITY:.0%} of "
              f"possible edges. A subject-state was analyzable if Neff ≥ "
              f"{MIN_NEFF}, the original network was valid, and ≥ "
              f"{MIN_SUCCESS:.0%} of resamples succeeded.")

    heading("Supplementary methods note: block-bootstrap stability")
    para(f"Model: {sel.LABEL}. Fit ID: {FIT_ID}.", size=8)
    para(methods_note(df), size=9)
    heading("Draft legend: Figure S# (FigS_bootstrap_attrition)")
    para(legend_text(df, sorted(df.state.unique())), size=9)

    add_table("Table S-A. Block-bootstrap stability of state-specific "
              "networks, by HMM state.", tA,
              "Median [IQR] across subject-states. † Analyzable subject-states "
              "only; the per-subject-state value is the median across "
              "resamples. Chance Jaccard at this density ≈ "
              f"{jaccard_chance(GRAPH_DENSITY):.3f}. " + common)
    add_table("Table S-B. Analyzable subject-states by group and state.", tB,
              "Analyzable/total (%). 'All states analyzable': subjects "
              "contributing an analyzable network for every state. "
              f"Near-static: subjects with a state FO ≥ {NEAR_STATIC_FO}. "
              + common)
    add_table("Table S-C. Subject-states excluded from stability analysis.",
              tC,
              f"Subject-states failing the occupancy criterion (Neff < "
              f"{MIN_NEFF}) were not bootstrapped. Resamples were counted as "
              f"unsuccessful when the resampled Neff fell below {MIN_NEFF} "
              "or the resampled network was invalid.")
    add_table("Table S-D. Stability by group and state "
              "(analyzable subject-states).", tD,
              "Median [IQR] across subjects within each group × state cell.")
    add_table("Table S-E. Group effects on network stability, with and "
              "without adjustment for Neff.", tE,
              "OLS with group and state as factors (reference group: "
              f"{group_order(df.group)[0]}); cluster-robust standard errors "
              "by subject, because states are nested within subjects. The "
              "Wald test is the joint test of all group terms. β for "
              "log Neff is the change in the stability metric per unit "
              "increase in log Neff. Analyzable subject-states only.")
    add_table("Table S-F. Near-static subject-states "
              f"(FO ≥ {NEAR_STATIC_FO}).", tF,
              "These subjects spent most of the scan in one state, so the "
              "state-specific network approximates static FC. 'Other states "
              "excluded' counts that subject's remaining subject-states that "
              "were not analyzable.")
    doc.save(path)


# ------------------------------------------------------------------
def main():
    global FIT_ID
    ap = argparse.ArgumentParser()
    ap.add_argument("--csv", default=str(sel.BOOTSTRAP_CSV))
    ap.add_argument("--out_dir", default=str(sel.MODEL_DIR / "bootstrap_supplement"))
    ap.add_argument("--n_roi", type=int, default=None)
    ap.add_argument("--min_neff", type=float, default=MIN_NEFF,
                    help="occupancy criterion used by run5B (hmm_config.m: 83)")
    args = ap.parse_args()
    os.makedirs(args.out_dir, exist_ok=True)
    set_min_neff(int(args.min_neff) if float(args.min_neff).is_integer() else args.min_neff)

    df = pd.read_csv(args.csv)
    FIT_ID = sel.check_fit_id(df, os.path.basename(args.csv)) or "not recorded"
    bools = ["passes_Neff", "valid_original_network", "analyzable",
             "split_half_valid", "near_static"]
    for c in bools:
        if c in df.columns:
            if df[c].dtype == object:       # "true"/"false" text: astype(bool) would
                df[c] = (df[c].astype(str).str.strip().str.lower()   # make all True
                         .isin(["true", "1", "1.0"]))
            else:
                df[c] = df[c].fillna(0).astype(bool)
    if "near_static" not in df.columns:
        df["near_static"] = df.FO >= NEAR_STATIC_FO
    df["group"] = df["group"].astype(str)
    df["status"] = df.apply(classify, axis=1)

    p = args.n_roi or int(round(np.nanmedian(df.p_over_Neff * df.Neff)))
    global N_BOOT
    pos = df[df.bootstrap_success_fraction > 0]
    if len(pos):
        N_BOOT = int(round(np.nanmedian(pos.bootstrap_success_count
                                        / pos.bootstrap_success_fraction)))
    states = sorted(df.state.unique().astype(int))
    groups = group_order(df.group)
    v2 = has(df, "median_edge_jaccard_rr")

    print(f"Subjects: {df.subject_id.nunique()}  States: {len(states)}  "
          f"p: {p}  CSV version: {'v2' if v2 else 'v1'}  Neff >= {MIN_NEFF}  "
          f"fit {FIT_ID[:12]}")
    print(df.groupby("state").status.value_counts().unstack(fill_value=0))

    fig_attrition(df, states, groups, args.out_dir)
    fig_stability(df, states, groups, args.out_dir)
    fig_occupancy(df, states, groups, args.out_dir, p)

    tables = build_tables(df, states, groups)
    for name, t in zip("ABCDEF", tables):
        if t is not None:
            t.to_csv(os.path.join(args.out_dir, f"TableS_{name}.csv"),
                     index=False)

    cov = df[["subject_id", "group", "state", "FO", "Neff", "near_static",
              "analyzable"]].copy()
    cov["fit_id"] = FIT_ID
    cov["log_Neff"] = np.log(cov.Neff.clip(lower=1e-12))
    cov.to_csv(os.path.join(args.out_dir, "neff_covariates.csv"),
               index=False)

    if tables[4] is None:
        print("statsmodels not installed: Table S-E skipped "
              "(pip install statsmodels).")
    else:
        with pd.option_context("display.width", 200,
                               "display.max_columns", 20):
            print(tables[4].to_string(index=False))
    try:
        write_docx(tables, os.path.join(args.out_dir,
                                        "TableS_bootstrap_stability.docx"),
                   df, p)
    except ImportError:
        print("python-docx not installed; wrote CSV tables only.")
    print(f"Done. Outputs in {args.out_dir}")


if __name__ == "__main__":
    main()